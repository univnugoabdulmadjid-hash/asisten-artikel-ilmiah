import streamlit as st
from google import genai
import tempfile
import os
import re
import time
import json
import urllib.request
import urllib.parse
from io import BytesIO
from docx import Document

# 1. KONFIGURASI HALAMAN
st.set_page_config(
    page_title="Asisten Riset & Publikasi Ilmiah Akademis",
    page_icon="🏛️",
    layout="wide",
    initial_sidebar_state="expanded"
)

st.markdown("""
    <style>
    .main { padding: 1rem; }
    .stButton>button { width: 100%; border-radius: 8px; height: 3em; font-weight: bold; }
    .stTextArea textarea { font-family: 'Times New Roman', serif; }
    </style>
""", unsafe_allow_html=True)

# PATH LOKAL REPOSITORY MEMORI LAPTOP YOGA
LOCAL_REPO_PATH = r"D:\KAMPUS UNUGO\02. TRIDHARMA PT\01. PUBLIKASI KARYA\01.ARTIKEL\01_memory_repository_webapp_madjid"

# HELPER KONVERSI MARKDOWN KE DOCX
def markdown_to_docx(md_text):
    doc = Document()
    lines = md_text.split('\n')
    for line in lines:
        line_s = line.strip()
        if not line_s:
            continue
        if line_s.startswith('# '):
            doc.add_heading(line_s[2:].replace('**', ''), level=1)
        elif line_s.startswith('## '):
            doc.add_heading(line_s[3:].replace('**', ''), level=2)
        elif line_s.startswith('### '):
            doc.add_heading(line_s[4:].replace('**', ''), level=3)
        elif line_s.startswith('#### '):
            doc.add_heading(line_s[5:].replace('**', ''), level=4)
        elif line_s.startswith('- ') or line_s.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            _add_formatted_text(p, line_s[2:])
        else:
            p = doc.add_paragraph()
            _add_formatted_text(p, line_s)
    
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

def _add_formatted_text(paragraph, text):
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)

# HELPER PEMINDAI BERKAS PROYEK LOKAL
def get_local_project_files():
    if os.path.exists(LOCAL_REPO_PATH):
        try:
            return [f for f in os.listdir(LOCAL_REPO_PATH) if f.endswith('.json')]
        except Exception:
            return []
    return []

# -------------------------------------------------------------
# HYBRID CITATION FETCHERS (CROSSREF + OPENALEX + GOOGLE BOOKS)
# -------------------------------------------------------------

def fetch_crossref_journals(query_topic, limit=20):
    try:
        clean_query = re.sub(r'[^\w\s]', '', query_topic)[:150]
        encoded_query = urllib.parse.quote(clean_query)
        url = f"https://api.crossref.org/works?query={encoded_query}&filter=type:journal-article&rows={limit}&sort=relevance"
        
        req = urllib.request.Request(
            url,
            headers={'User-Agent': 'AsistenRisetAkademis/1.0 (mailto:abdulmadjidpodungge@unugo.ac.id)'}
        )
        
        with urllib.request.urlopen(req, timeout=10) as response:
            data = json.loads(response.read().decode())
            items = data.get('message', {}).get('items', [])
            
            citations = []
            for item in items:
                title = item.get('title', [''])[0]
                authors_raw = item.get('author', [])
                author_list = []
                for a in authors_raw:
                    family = a.get('family', '')
                    given = a.get('given', '')
                    if family:
                        author_list.append({'family': family, 'given': given, 'formatted': f"{family}, {given[0] if given else ''}.".strip()})
                
                authors_str = ", ".join([a['formatted'] for a in author_list]) if author_list else "Anonim"
                
                pub_date = item.get('published-print') or item.get('published-online') or {}
                date_parts = pub_date.get('date-parts', [[None]])[0]
                year = date_parts[0] if date_parts and date_parts[0] else "n.d."
                
                journal = item.get('container-title', [''])[0]
                volume = item.get('volume', '')
                issue = item.get('issue', '')
                doi = item.get('DOI', '')
                
                if title and doi:
                    ref_apa = f"{authors_str} ({year}). {title}. {journal}"
                    if volume:
                        ref_apa += f", {volume}"
                    if issue:
                        ref_apa += f"({issue})"
                    ref_apa += f". https://doi.org/{doi}"
                    
                    citations.append({
                        'type': 'JOUR',
                        'apa': ref_apa,
                        'doi': doi,
                        'authors': author_list,
                        'authors_str': authors_str,
                        'year': year,
                        'title': title,
                        'journal': journal,
                        'volume': volume,
                        'issue': issue
                    })
            return citations[:20]
    except Exception:
        return []

def fetch_openalex_books(query_topic, limit=10):
    try:
        clean_query = re.sub(r'[^\w\s]', '', query_topic)[:150]
        encoded_query = urllib.parse.quote(clean_query)
        url = f"https://api.openalex.org/works?search={encoded_query}&filter=type:book|book-chapter&per_page={limit}"
        
        req = urllib.request.Request(
            url,
            headers={'User-Agent': 'AsistenRisetAkademis/1.0 (mailto:abdulmadjidpodungge@unugo.ac.id)'}
        )
        
        with urllib.request.urlopen(req, timeout=10) as response:
            data = json.loads(response.read().decode())
            results = data.get('results', [])
            
            books = []
            for item in results:
                title = item.get('title', '')
                authorships = item.get('authorships', [])
                author_list = []
                for auth in authorships:
                    name = auth.get('author', {}).get('display_name', '')
                    if name:
                        parts = name.split(' ')
                        family = parts[-1] if len(parts) > 1 else name
                        given = " ".join(parts[:-1]) if len(parts) > 1 else ""
                        author_list.append({'family': family, 'given': given, 'formatted': f"{family}, {given[0] if given else ''}.".strip()})
                
                authors_str = ", ".join([a['formatted'] for a in author_list]) if author_list else "Anonim"
                year = item.get('publication_year', 'n.d.')
                doi_raw = item.get('doi', '')
                doi = doi_raw.replace('https://doi.org/', '') if doi_raw else ''
                publisher = item.get('primary_location', {}).get('source', {}).get('display_name', 'Penerbit Akademis')
                
                if title:
                    ref_apa = f"{authors_str} ({year}). *{title}*. {publisher}."
                    if doi:
                        ref_apa += f" https://doi.org/{doi}"
                    
                    books.append({
                        'type': 'BOOK',
                        'apa': ref_apa,
                        'doi': doi,
                        'authors': author_list,
                        'authors_str': authors_str,
                        'year': year,
                        'title': title,
                        'publisher': publisher
                    })
            return books
    except Exception:
        return []

def fetch_google_books(query_topic, limit=10):
    try:
        clean_query = re.sub(r'[^\w\s]', '', query_topic)[:150]
        encoded_query = urllib.parse.quote(clean_query)
        url = f"https://www.googleapis.com/books/v1/volumes?q={encoded_query}&maxResults={limit}"
        
        req = urllib.request.Request(url)
        with urllib.request.urlopen(req, timeout=10) as response:
            data = json.loads(response.read().decode())
            items = data.get('items', [])
            
            books = []
            for item in items:
                vol_info = item.get('volumeInfo', {})
                title = vol_info.get('title', '')
                authors_raw = vol_info.get('authors', [])
                author_list = []
                for name in authors_raw:
                    parts = name.split(' ')
                    family = parts[-1] if len(parts) > 1 else name
                    given = " ".join(parts[:-1]) if len(parts) > 1 else ""
                    author_list.append({'family': family, 'given': given, 'formatted': f"{family}, {given[0] if given else ''}.".strip()})
                
                authors_str = ", ".join([a['formatted'] for a in author_list]) if author_list else "Anonim"
                pub_date = vol_info.get('publishedDate', 'n.d.')
                year = pub_date[:4] if len(pub_date) >= 4 else 'n.d.'
                publisher = vol_info.get('publisher', 'Penerbit Referensi')
                
                industry_ids = vol_info.get('industryIdentifiers', [])
                isbn = ""
                for ind in industry_ids:
                    if ind.get('type') in ['ISBN_13', 'ISBN_10']:
                        isbn = ind.get('identifier', '')
                        break
                
                if title:
                    ref_apa = f"{authors_str} ({year}). *{title}*. {publisher}."
                    if isbn:
                        ref_apa += f" ISBN: {isbn}."
                    
                    books.append({
                        'type': 'BOOK',
                        'apa': ref_apa,
                        'doi': f"ISBN-{isbn}" if isbn else "",
                        'authors': author_list,
                        'authors_str': authors_str,
                        'year': year,
                        'title': title,
                        'publisher': publisher
                    })
            return books
    except Exception:
        return []

# HELPER KONVERSI SITASI KE KODE RIS (MENDELEY / ZOTERO)
def convert_citations_to_ris(citations):
    ris_output = ""
    for item in citations:
        if isinstance(item, dict):
            c_type = item.get('type', 'JOUR')
            ris_output += f"TY  - {c_type}\n"
            if item.get('title'):
                ris_output += f"TI  - {item.get('title')}\n"
            if item.get('journal'):
                ris_output += f"JO  - {item.get('journal')}\n"
            if item.get('publisher'):
                ris_output += f"PB  - {item.get('publisher')}\n"
            if item.get('year'):
                ris_output += f"PY  - {item.get('year')}\n"
            if item.get('volume'):
                ris_output += f"VL  - {item.get('volume')}\n"
            if item.get('issue'):
                ris_output += f"IS  - {item.get('issue')}\n"
            if item.get('doi'):
                ris_output += f"DO  - {item.get('doi')}\n"
                ris_output += f"UR  - https://doi.org/{item.get('doi')}\n"
            for auth in item.get('authors', []):
                if isinstance(auth, dict):
                    ris_output += f"AU  - {auth.get('family', '')}, {auth.get('given', '')}\n"
            ris_output += "ER  - \n\n"
    return ris_output

# HELPER PENYANGGA OTOMATIS (AUTO RETRY & MODEL FALLBACK)
def generate_content_with_retry(client, primary_model, contents):
    fallback_sequence = [primary_model, "gemini-1.5-flash", "gemini-1.5-pro", "gemini-2.5-flash"]
    models_to_try = list(dict.fromkeys(fallback_sequence))
    
    last_error = None
    for model_name in models_to_try:
        for attempt in range(2):
            try:
                response = client.models.generate_content(
                    model=model_name,
                    contents=contents
                )
                return response
            except Exception as e:
                last_error = e
                err_msg = str(e)
                if "503" in err_msg or "UNAVAILABLE" in err_msg or "429" in err_msg:
                    time.sleep(2)
                    continue
                else:
                    break
    raise last_error

# 2. PANEL SIDEBAR & PROTEKSI PIN
st.sidebar.title("⚙️ Pengaturan Sistem")

if 'authenticated' not in st.session_state:
    st.session_state['authenticated'] = False

for key in ['synthesis_result', 'outline_result', 'draft_result', 'revision_matrix', 'hybrid_citations', 'project_name', 'project_notes']:
    if key not in st.session_state:
        st.session_state[key] = None

app_pin = st.secrets.get("APP_PIN", "")

if app_pin:
    if not st.session_state['authenticated']:
        user_pin = st.sidebar.text_input("Masukkan PIN Akses Aplikasi:", type="password", help="Masukkan kata sandi/PIN akses")
        if user_pin == app_pin:
            st.session_state['authenticated'] = True
            st.rerun()
        elif user_pin != "":
            st.sidebar.error("❌ PIN yang Anda masukkan salah!")
            
    if not st.session_state['authenticated']:
        st.warning("🔒 **Aplikasi Terkunci.** Masukkan PIN Akses yang benar pada sidebar untuk membuka seluruh fitur.")
        st.stop()
    else:
        if st.sidebar.button("🔒 Logout / Kunci Aplikasi"):
            st.session_state['authenticated'] = False
            st.session_state.clear()
            st.rerun()

api_key = st.secrets.get("GEMINI_API_KEY", "")

selected_model = st.sidebar.selectbox(
    "Pilih Model Utama:",
    ["gemini-2.5-flash", "gemini-1.5-flash", "gemini-1.5-pro"],
    index=0
)

st.sidebar.markdown("---")

# MANAJEMEN PROYEK ARTIKEL
st.sidebar.subheader("📂 MANAJEMEN PROYEK ARTIKEL")

if st.sidebar.button("➕ Proyek Baru (Reset Sesi)", help="Bersihkan seluruh data draf saat ini untuk mulai analisis artikel baru"):
    for key in ['synthesis_result', 'outline_result', 'draft_result', 'revision_matrix', 'hybrid_citations', 'project_name', 'project_notes']:
        st.session_state[key] = None
    st.sidebar.success("Sesi berhasil dibersihkan. Siap untuk proyek baru!")
    st.rerun()

st.sidebar.markdown("---")

local_files = get_local_project_files()
if local_files:
    selected_local_file = st.sidebar.selectbox("📂 Pilih Proyek dari Folder Repository:", ["-- Pilih Proyek --"] + local_files)
    if selected_local_file != "-- Pilih Proyek --":
        if st.sidebar.button("📥 Buka Proyek Terpilih"):
            try:
                full_path = os.path.join(LOCAL_REPO_PATH, selected_local_file)
                with open(full_path, "r", encoding="utf-8") as f:
                    loaded_data = json.load(f)
                st.session_state['project_name'] = loaded_data.get('project_name')
                st.session_state['project_notes'] = loaded_data.get('project_notes')
                st.session_state['synthesis_result'] = loaded_data.get('synthesis_result')
                st.session_state['outline_result'] = loaded_data.get('outline_result')
                st.session_state['draft_result'] = loaded_data.get('draft_result')
                st.session_state['revision_matrix'] = loaded_data.get('revision_matrix')
                st.session_state['hybrid_citations'] = loaded_data.get('hybrid_citations')
                st.sidebar.success(f"✅ Proyek '{selected_local_file}' berhasil dimuat!")
                st.rerun()
            except Exception as e:
                st.sidebar.error(f"Gagal memuat berkas: {e}")
else:
    st.sidebar.info("💡 Belum ada proyek tersimpan di folder lokal.")

uploaded_project = st.sidebar.file_uploader("📂 Muat Berkas Proyek (.json):", type=["json"], key="project_loader")
if uploaded_project is not None:
    if st.sidebar.button("📥 Pulihkan dari File Upload"):
        try:
            loaded_data = json.load(uploaded_project)
            st.session_state['project_name'] = loaded_data.get('project_name')
            st.session_state['project_notes'] = loaded_data.get('project_notes')
            st.session_state['synthesis_result'] = loaded_data.get('synthesis_result')
            st.session_state['outline_result'] = loaded_data.get('outline_result')
            st.session_state['draft_result'] = loaded_data.get('draft_result')
            st.session_state['revision_matrix'] = loaded_data.get('revision_matrix')
            st.session_state['hybrid_citations'] = loaded_data.get('hybrid_citations')
            st.sidebar.success("✅ Proyek berhasil dipulihkan!")
            st.rerun()
        except Exception as e:
            st.sidebar.error("Gagal membaca berkas proyek.")

st.sidebar.markdown("---")
st.sidebar.markdown("### 💾 Simpan / Update Proyek")

proj_name_val = st.sidebar.text_input("Nama Artikel / ID Proyek (*Wajib):", value=st.session_state['project_name'] or "Draf_Artikel_Hukum_HAN")
st.session_state['project_name'] = proj_name_val

proj_notes_val = st.sidebar.text_area("Catatan Khusus Penulis (Opsional):", value=st.session_state['project_notes'] or "", placeholder="Misal: Target Jurnal SINTA 2 PJIH, fokus asas kepastian hukum...")
st.session_state['project_notes'] = proj_notes_val

project_export_data = {
    'project_name': st.session_state['project_name'],
    'project_notes': st.session_state['project_notes'],
    'synthesis_result': st.session_state['synthesis_result'],
    'outline_result': st.session_state['outline_result'],
    'draft_result': st.session_state['draft_result'],
    'revision_matrix': st.session_state['revision_matrix'],
    'hybrid_citations': st.session_state['hybrid_citations']
}
json_project_str = json.dumps(project_export_data, indent=2)
clean_proj_filename = re.sub(r'[^\w\s-]', '', st.session_state['project_name']).strip().replace(' ', '_')

if os.path.exists(LOCAL_REPO_PATH):
    if st.sidebar.button("💾 Simpan ke Repository Lokal"):
        try:
            target_file_path = os.path.join(LOCAL_REPO_PATH, f"{clean_proj_filename}_Session.json")
            with open(target_file_path, "w", encoding="utf-8") as f:
                f.write(json_project_str)
            st.sidebar.success(f"✅ Proyek berhasil disimpan ke:\n`01_memory_repository_webapp_madjid`")
            time.sleep(1)
            st.rerun()
        except Exception as e:
            st.sidebar.error(f"Gagal menyimpan ke folder lokal: {e}")

st.sidebar.download_button(
    label="📥 Unduh Sesi Proyek (.json)",
    data=json_project_str,
    file_name=f"{clean_proj_filename}_Session.json",
    mime="application/json",
    help="Unduh berkas JSON proyek untuk cadangan manual."
)

st.sidebar.markdown("---")
st.sidebar.subheader("🎯 Domain Keilmuan Riset")
domain_mode = st.sidebar.radio(
    "Pilih Kategori Skema Riset:",
    ["Mode Hukum (Utama)", "Mode Multidisiplin"]
)

sub_discipline = ""
data_type_mode = "Kualitatif / Deskriptif / Doktrinal"

if domain_mode == "Mode Hukum (Utama)":
    sub_discipline = st.sidebar.selectbox(
        "Pilih Cabang / Sub-Kepakaran Hukum:",
        [
            "HAN, Perlindungan Hukum, Perbandingan & Filsafat Hukum (Kepakaran Utama)",
            "Hukum Pidana",
            "Hukum Perdata",
            "Hukum Tata Negara & Konstitusi",
            "Hukum Internasional & Cabang Hukum Lainnya"
        ]
    )
else:
    sub_discipline = st.sidebar.text_input(
        "Sebutkan Bidang Keilmuan / Ranting Riset:",
        value="Teknik Arsitektur / Administrasi Publik / Kesejahteraan Sosial",
        help="Contoh: Teknik Arsitektur, Teknik Lingkungan, Administrasi Publik, Pendidikan, Keagamaan Islam, dll."
    )
    data_type_mode = st.sidebar.radio(
        "Karakteristik Data Utama Riset:",
        [
            "Data Kualitatif / Deskriptif / Naratif / Dokumen",
            "Data Eksakta / Uji Eksperimen / Sampling Lab / Spesifikasi Maket"
        ]
    )

st.sidebar.markdown("---")
st.sidebar.info("📱 **Akses Lintas Perangkat**\nAplikasi siap dibuka melalui browser di Samsung S26 Ultra maupun TAB Huawei 12x.")

# ATURAN EMAS SISTEM (GLOBAL MANDATORY RULES)
COMMON_GOLDEN_RULES = """
# ATURAN EMAS PENULISAN ILMIAH (UNIVERSAL MANDATORY RULES)
1. ZERO HALLUCINATION (NOL HALUSINASI): Dilarang keras membuat citasi fiktif, merubah data uji, atau mengarang referensi. Seluruh rujukan wajib nyata dan memiliki identitas DOI / ISBN / Metadata resmi.
2. ATURAN KETAT JUDUL ARTIKEL:
   - Panjang Judul: MAKSIMAL 12 KATA (lugas, padat, berbobot).
   - Tanda Baca Terlarang: DILARANG KERAS menggunakan tanda titik dua (:).
   - Pembersihan Frasa: Buang frasa administratif/seremonial (seperti 'Laporan Pengabdian...', 'Proposal Hibah BIMA...').
3. KETENTUAN ABSTRAK & KATA KUNCI:
   - Abstrak: 150 hingga 250 kata dalam 1 paragraf utuh (memuat latar belakang singkat, tujuan, metode, hasil utama, dan implikasi/novelty).
   - Kata Kunci (Keywords): Terdiri dari 3 hingga 5 kata kunci utama, WAJIB DISUSUN SECARA ALFABETIS (A-Z), dan dipisahkan tanda titik koma (;).
4. KETENTUAN RUMUSAN MASALAH: Tepat menetapkan 2 Rumusan Masalah dalam bentuk kalimat tanya yang tajam dan fungsional.
5. KETENTUAN BAB HASIL DAN PEMBAHASAN:
   - Bab Hasil: Sintesis fenomena/data uji dengan literatur. Sekurang-kurangnya terdapat 3 Sub-Sub Bab Kunci.
   - Bab Pembahasan: Kontribusi ilmiah sangat unik (Novelty) terkait kepakaran. Sekurang-kurangnya terdapat 3 Sub-Sub Bab Kunci.
   - Bagian Wajib (Limitation & Future Research): Wajib dicantumkan di akhir Bab Pembahasan untuk membuktikan posisi naskah sebagai extend atau challenge.
6. LARANGAN STRUKTUR KALIMAT ANTITESIS AI: DILARANG KERAS menggunakan pola kalimat antitesis berulang seperti 'tidak hanya X, tetapi juga Y' atau 'ini bukan hanya... melainkan...'. Gunakan pernyataan langsung, ekspresif, dan alami (human-like).
7. STRUKTUR DAFTAR PUSTAKA KATEGORIS (AKADEMIK HUKUM BAKU):
   - Daftar Pustaka WAJIB dipisahkan menjadi 2 KELOMPOK SUB-JUDUL UTAMA:
     a. Buku / Monograf / Book Chapter
     b. Jurnal Ilmiah / Artikel Jurnal
   - Setiap kelompok rujukan WAJIB DISUSUN SECARA ALFABETIS dari A sampai Z (A-Z) berformat APA Style 7th Edition.
8. KOMPOSISI RUJUKAN HYBRID:
   - Wajib menyerap minimal 8 RUJUKAN BUKU / BOOK CHAPTER / MONOGRAF resmi.
   - Wajib menyerap minimal 15 hingga 20 ARTIKEL JURNAL resmi ber-DOI.
9. MAKSIMAL PANJANG NASKAH: Keseluruhan draf naskah maksimal 4.000 kata (sudah termasuk Daftar Pustaka berformat APA Style 7th Edition).
"""

if domain_mode == "Mode Hukum (Utama)":
    SYSTEM_INSTRUCTION = COMMON_GOLDEN_RULES + f"""
    # PERAN DAN IDENTITAS SISTEM
    Anda adalah Asisten Riset Hukum Senior dan Co-Author Akademis spesialisasi {sub_discipline}.

    # KETENTUAN HUKUM SPESIFIK
    - DISIPLIN ILMU: Murni doktrin hukum, teoretis, yuridis normatif/empiris. DILARANG KERAS mengarahkan ke Administrasi Publik atau Pelayanan Publik umum.
    - INTEGRASI ADAGIUM & TEORI HUKUM: Sisipkan adagium hukum Latin/azas hukum relevan serta teori hukum utama yang presisi.
    - FORMAT CITASI WAJIB: APA Style 7th Edition.
    """
else:
    SYSTEM_INSTRUCTION = COMMON_GOLDEN_RULES + f"""
    # PERAN DAN IDENTITAS SISTEM
    Anda adalah Asisten Riset Multidisiplin Senior dan Co-Author Akademis khusus mendampingi Dosen/Peneliti pada bidang keilmuan: {sub_discipline}.

    # KETENTUAN MULTIDISIPLIN SPESIFIK
    - AKURASI DATA EKSAKTA / SAMPLING: Jika disediakan data hasil pengujian laboratorium/maket/sampel, gunakan angka dan parameter tersebut secara mutlak tanpa rekayasa.
    - FORMAT CITASI WAJIB: APA Style 7th Edition.
    """

active_proj_display = st.session_state['project_name'] if st.session_state['project_name'] else "Proyek Baru (Unsaved)"
st.title("🏛️ Aplikasi Asisten Penulisan Artikel Ilmiah & Revisi Jurnal")
st.info(f"📌 **Proyek Aktif:** `{active_proj_display}` | **Status Sistem:** `{domain_mode}` ({sub_discipline})")

if not api_key:
    st.warning("⚠️ Kunci API belum terdeteksi pada Secrets sistem.")
    st.stop()

try:
    client = genai.Client(api_key=api_key.strip())
except Exception as e:
    st.error(f"Gagal menginisialisasi API Key: {e}")
    st.stop()

# 3. NAVIGASI ALUR KERJA SEKUANSIAL (5 TABS)
tab1, tab2, tab3, tab4, tab5 = st.tabs([
    "1️⃣ Fitur 1-3: Sintesis & SOTA", 
    "2️⃣ Fitur 4-5: Data Riset & Outline", 
    "3️⃣ Fitur 6-7: Draf Naskah & Ekspor",
    "4️⃣ Asisten Revisi Jurnal (Peer Review)",
    "5️⃣ Ekspor Sitasi RIS (Mendeley)"
])

# TAB 1: SINTESIS LITERATUR & SOTA
with tab1:
    st.header("Analisis Literatur & Rekomendasi Judul")
    mode_input = st.radio(
        "Pilih Metode Input Data:", 
        [
            "Unggah Berkas PDF Literasi / Rujukan (Fitur 1)", 
            "Input Manual Tabel Sintesis (Fitur 2)",
            "Ekstraksi Naskah Proposal & Laporan Riset/Pengabdian (Fitur 3)"
        ]
    )
    
    if mode_input == "Unggah Berkas PDF Literasi / Rujukan (Fitur 1)":
        uploaded_files = st.file_uploader("Unggah Artikel Ilmiah / Jurnal Rujukan (PDF):", type=["pdf"], accept_multiple_files=True)
        if st.button("🚀 Ekstrak & Buat Sintesis (3 Rekomendasi Judul)", key="btn_synthesis_pdf"):
            if uploaded_files:
                with st.spinner("AI sedang mengunggah & mengindeks seluruh berkas PDF ke server..."):
                    try:
                        temp_files = []
                        uploaded_gemini_files = []
                        
                        for file in uploaded_files:
                            with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
                                tmp.write(file.getvalue())
                                temp_files.append(tmp.name)
                            gfile = client.files.upload(file=tmp.name)
                            uploaded_gemini_files.append(gfile)
                        
                        status_box = st.info("Memastikan seluruh dokumen selesai terindeks secara utuh...")
                        for gfile in uploaded_gemini_files:
                            while gfile.state.name == "PROCESSING":
                                time.sleep(2)
                                gfile = client.files.get(name=gfile.name)
                            if gfile.state.name == "FAILED":
                                raise Exception(f"Gagal memproses berkas: {gfile.display_name}")
                        status_box.empty()
                        
                        with st.spinner("Menganalisis literatur, menyusun 5 SOTA, dan 3 Rekomendasi Judul..."):
                            prompt = SYSTEM_INSTRUCTION + """\n
                            Tugas: Ekstrak seluruh isi PDF rujukan dan hasilkan 3 REKOMENDASI JUDUL UTAMA.
                            
                            Persyaratan Judul:
                            - MAKSIMAL 12 KATA PER JUDUL.
                            - DILARANG KERAS MENGGUNAKAN TANDA TITIK DUA (:).
                            
                            Format Keluaran Setiap Rekomendasi:
                            - Judul (Maksimal 12 kata, tanpa tanda titik dua)
                            - Metodologi / Pendekatan Riset
                            - Gap Research (Analisis mendalam celah penelitian)
                            - 5 SOTA (Kutipan/sintesis ilmiah dalam format APA Style 7th Edition)
                            - Novelty (Kontribusi ilmiah paling unik dan berbobot)
                            """
                            response = generate_content_with_retry(
                                client,
                                selected_model,
                                [prompt] + uploaded_gemini_files
                            )
                            st.session_state['synthesis_result'] = response.text
                        
                        for path in temp_files:
                            if os.path.exists(path):
                                os.remove(path)
                    except Exception as e:
                        st.error(f"Terjadi kesalahan pemrosesan: {e}")
            else:
                st.warning("Mohon unggah setidaknya 1 file PDF.")
                
    elif mode_input == "Input Manual Tabel Sintesis (Fitur 2)":
        manual_text = st.text_area("Masukkan tabel/informasi sintesis Anda di sini:", height=200)
        if st.button("🚀 Buat Sintesis dari Input Manual", key="btn_synthesis_manual"):
            if manual_text:
                with st.spinner("Memproses sintesis manual..."):
                    prompt = SYSTEM_INSTRUCTION + f"""\n
                    Tugas: Sintesiskan informasi manual berikut menjadi 3 REKOMENDASI JUDUL UTAMA.
                    - Masing-masing judul MAKSIMAL 12 KATA dan DILARANG KERAS menggunakan tanda titik dua (:).
                    - Lengkapi dengan Gap Research, 5 SOTA (APA 7th), dan Novelty.
                    \nData Input:\n{manual_text}
                    """
                    response = generate_content_with_retry(
                        client,
                        selected_model,
                        prompt
                    )
                    st.session_state['synthesis_result'] = response.text
            else:
                st.warning("Mohon isi teks sintesis terlebih dahulu.")

    else: # Fitur 3: Ekstraksi Naskah Proposal & Laporan Riset/Pengabdian
        st.subheader("📄 Ekstraksi Naskah Proposal & Laporan Riset/Pengabdian")
        doc_type = st.radio(
            "Pilih Jenis Dokumen yang Diunggah:",
            ["Proposal Penelitian (e.g. Hibah BIMA / DIPA)", "Laporan Pengabdian kepada Masyarakat (PkM)"]
        )
        prop_file = st.file_uploader("Unggah Berkas Proposal / Laporan (PDF):", type=["pdf"], key="uploader_prop")
        
        if st.button("🚀 Ekstrak & Buat Sintesis (4 Opsi Judul)", key="btn_synthesis_prop"):
            if prop_file:
                with st.spinner("AI sedang mengunggah & menganalisis berkas proposal/laporan..."):
                    try:
                        with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
                            tmp.write(prop_file.getvalue())
                            tmp_path = tmp.name
                        
                        gfile = client.files.upload(file=tmp_path)
                        while gfile.state.name == "PROCESSING":
                            time.sleep(2)
                            gfile = client.files.get(name=gfile.name)
                        
                        with st.spinner("Mengekstrak substansi, menyusun 5 SOTA, dan 4 Opsi Judul Jurnal..."):
                            prompt_prop = SYSTEM_INSTRUCTION + f"""\n
                            Tugas: Ekstrak seluruh isi dokumen ({doc_type}) dan hasilkan 4 OPSI REKOMENDASI JUDUL ARTIKEL JURNAL.

                            PERSYARATAN KETAT UNTUK SEMUA JUDUL:
                            1. MAKSIMAL 12 KATA PER JUDUL.
                            2. DILARANG KERAS MENGGUNAKAN TANDA TITIK DUA (:).
                            3. HAPUS SELURUH FRASA ADMINISTRATIF / SEREMONIAL.

                            PEMBAGIAN 4 OPSI JUDUL:
                            - OPSI 1 (Judul Asli Refined): Dipoles dari judul asli dokumen agar memenuhi standar jurnal bereputasi (maksimal 12 kata, tanpa titik dua).
                            - OPSI 2 (Variasi Tematik A): Penekanan pada sudut pandang yuridis/doktrinal atau substansi utama.
                            - OPSI 3 (Variasi Tematik B): Penekanan pada sudut pandang efektivitas prosedural, instrumen perlindungan, atau dampak pelaksanaan.
                            - OPSI 4 (Variasi Tematik C): Penekanan pada sudut pandang filsafat hukum, asas hukum, atau perbandingan.

                            FORMAT KELUARAN UNTUK SETIAP OPSI JUDUL:
                            - Judul Artikel (Maksimal 12 kata, tanpa titik dua)
                            - Metodologi / Pendekatan Riset
                            - Gap Research / Analisis Dampak Permasalahan
                            - 5 SOTA Rujukan Pendukung (APA Style 7th Edition)
                            - Novelty (Kontribusi ilmiah paling unik)
                            """
                            response = generate_content_with_retry(
                                client,
                                selected_model,
                                [prompt_prop, gfile]
                            )
                            st.session_state['synthesis_result'] = response.text
                        
                        if os.path.exists(tmp_path):
                            os.remove(tmp_path)
                    except Exception as e:
                        st.error(f"Gagal memproses dokumen: {e}")
            else:
                st.warning("Mohon unggah berkas PDF proposal/laporan terlebih dahulu.")

    if st.session_state['synthesis_result']:
        st.markdown("---")
        st.subheader("📊 Hasil Matriks Sintesis & Rekomendasi Judul")
        st.markdown(st.session_state['synthesis_result'])

# TAB 2: DATA RISET & OUTLINE
with tab2:
    st.header("Klasifikasi Metodologi & Penyusunan Outline")
    
    if not st.session_state['synthesis_result']:
        st.warning("🔒 **Tahap Terkunci.** Silakan selesaikan Tahap 1 (Sintesis & SOTA) terlebih dahulu.")
    else:
        opsi_judul = st.selectbox("Pilih Opsi Rekomendasi Judul yang Ingin Dilanjutkan:", [
            "Opsi / Rekomendasi 1",
            "Opsi / Rekomendasi 2",
            "Opsi / Rekomendasi 3",
            "Opsi / Rekomendasi 4 (Khusus Fitur 3)"
        ])
        
        if "Eksakta" in data_type_mode:
            st.warning("🧪 **Mode Data Eksakta / Uji Laboratorium Aktif**")
            data_input_field = st.text_area(
                "Masukkan Data Hasil Uji Eksperimen / Sampling Laboratorium / Spesifikasi Maket / Parameter Fisik:",
                height=180,
                placeholder="Contoh: Hasil uji kuat tekan beton (MPa), efisiensi ventilasi alami (%), kadar pencemaran air sampel A & B, spesifikasi material maket..."
            )
        else:
            data_input_field = st.text_area(
                "Masukkan Data Lapangan / Data Empiris / Hasil Wawancara / Dokumen Riset Tambahan (Opsional):",
                height=150,
                placeholder="Contoh: Hasil wawancara, data statistik kasus, dokumen pemeriksaan..."
            )
            
        if st.button("📑 Buat Rekomendasi Outline Naskah"):
            with st.spinner("AI sedang merancang kerangka penulisan..."):
                prompt_outline = SYSTEM_INSTRUCTION + f"""\n
                Berdasarkan sintesis sebelumnya:
                {st.session_state['synthesis_result']}
                
                Pilihan Pengguna: {opsi_judul}
                Data Tambahan: {data_input_field if data_input_field else 'Tidak ada (Pendekatan Teoretis/Kualitatif)'}
                
                Tugas: Hasilkan Outline Naskah Terstruktur yang WAJIB MEMATUHI ATURAN AKADEMIS berikut:
                - Judul: Maksimal 12 kata, TANPA tanda titik dua (:).
                - Abstrak: Draf kerangka abstrak 150 - 250 kata dalam 1 paragraf utuh.
                - Kata Kunci (Keywords): 3 hingga 5 kata kunci utama, WAJIB ALFABETIS (A-Z), dipisah tanda titik koma (;).
                - A. Pendahuluan (Sertakan fenomena, gap research, SOTA, novelty, teori utama, dan TEPAT 2 Rumusan Masalah bentuk kalimat tanya).
                - B. Metode Penelitian (Jenis penelitian & skala wilayah).
                - C. Hasil dan Pembahasan:
                    * Sub-Bab Hasil: Sekurang-kurangnya 3 Sub-Sub Bab Kunci.
                    * Sub-Bab Pembahasan: Sekurang-kurangnya 3 Sub-Sub Bab Kunci (Novelty).
                    * Bagian Wajib: Limitation & Future Research.
                - D. Kesimpulan (Menjawab 2 rumusan masalah secara langsung).
                - E. Daftar Pustaka (Wajib dikelompokkan: Buku A-Z dan Jurnal A-Z berformat APA Style 7th Edition).
                """
                response = generate_content_with_retry(
                    client,
                    selected_model,
                    [prompt_outline]
                )
                st.session_state['outline_result'] = response.text

    if st.session_state['outline_result']:
        st.markdown("---")
        st.subheader("📌 Rekomendasi Outline Naskah")
        st.markdown(st.session_state['outline_result'])
        
        st.markdown("---")
        st.subheader("💬 Ruang Diskusi & Revisi Outline")
        revisi_input = st.text_input("Catatan Revisi Outline (opsional):", placeholder="Misal: Penjelasannya tolong lebih menonjolkan analisis asas kepastian hukum...")
        if st.button("🔄 Perbarui Outline Sesuai Catatan"):
            with st.spinner("Memperbarui outline..."):
                prompt_revisi = f"Berikut outline saat ini:\n{st.session_state['outline_result']}\n\nPermintaan Revisi Pengguna:\n{revisi_input}\n\nTolong perbarui outline tersebut dengan tetap mematuhi seluruh Aturan Emas."
                response = generate_content_with_retry(
                    client,
                    selected_model,
                    prompt_revisi
                )
                st.session_state['outline_result'] = response.text
                st.rerun()

# TAB 3: DRAF NASKAH & EKSPOR DOCX (KATEGORISASI DAFTAR PUSTAKA A-Z)
with tab3:
    st.header("Penulisan Draf Naskah Lengkap & Ekspor")
    
    if not st.session_state['outline_result']:
        st.warning("🔒 **Tahap Terkunci.** Silakan selesaikan dan setujui Outline pada Tab 2 terlebih dahulu.")
    else:
        st.success("Outline disetujui. Siap menyusun draf naskah dengan rujukan Hybrid & Daftar Pustaka Kategori A-Z.")
        
        if st.button("✍️ Generasi Draf Naskah Lengkap (Maksimal 4.000 Kata)"):
            outline_text = st.session_state['outline_result']
            
            with st.spinner("🔍 Menghubungi Crossref, OpenAlex, dan Google Books API untuk menarik rujukan BUKU & JURNAL mutakhir..."):
                journals = fetch_crossref_journals(outline_text[:150], limit=20)
                openalex_bks = fetch_openalex_books(outline_text[:150], limit=5)
                gbooks = fetch_google_books(outline_text[:150], limit=5)
                
                all_books = openalex_bks + gbooks
                all_citations = journals + all_books
                st.session_state['hybrid_citations'] = all_citations
                
                formatted_refs_str = ""
                formatted_refs_str += "=== KATEGORI RUJUKAN BUKU / MONOGRAF / BOOK CHAPTER (MINIMAL 8) ===\n"
                for b in all_books:
                    formatted_refs_str += f"- [BUKU] {b['apa']}\n"
                    
                formatted_refs_str += "\n=== KATEGORI RUJUKAN ARTIKEL JURNAL DOI (MINIMAL 15-20) ===\n"
                for j in journals:
                    formatted_refs_str += f"- [JURNAL] {j['apa']}\n"

            with st.spinner("AI sedang menyusun naskah akademik utuh (Zero Hallucination, APA 7th, Anti-AI Detector)..."):
                prompt_draft = SYSTEM_INSTRUCTION + f"""\n
                Susun naskah artikel ilmiah lengkap secara utuh dan terstruktur berdasarkan Outline berikut:
                {outline_text}
                
                DAFTAR RUJUKAN RESMI HYBRID (WAJIB DISERAP SECARA KONTEKSTUAL KE PARAGRAF NASKAH):
                {formatted_refs_str}
                
                PETUNJUK KETAT PENULISAN NASKAH LENGKAP & DAFTAR PUSTAKA:
                1. Judul: Maksimal 12 kata, DILARANG KERAS tanda titik dua (:).
                2. Abstrak: Presisi 150 - 250 kata dalam 1 paragraf utuh.
                3. Kata Kunci (Keywords): 3 - 5 kata kunci, WAJIB ALFABETIS (A-Z), dipisah tanda titik koma (;).
                4. Panjang Naskah: Maksimal 4.000 kata komprehensif.
                5. PENGGUNAAN SITASI HYBRID: Wajib menyerap MINIMAL 8 BUKU/MONOGRAF dan 15-20 ARTIKEL JURNAL di atas ke dalam paragraf yang relevan secara kontekstual.
                6. Bahasa Akademis Formal Human-Like. DILARANG STRUKTUR ANTITESIS ('tidak hanya X tapi Y').
                7. Tepat menjawab 2 Rumusan Masalah.
                8. Bab Hasil (3 Sub-Sub Bab) & Bab Pembahasan (3 Sub-Sub Bab + Limitation & Future Research).
                9. FORMAT DAFTAR PUSTAKA WAJIB DIBAGI DUA KELOMPOK DENGAN URUTAN ALFABETIS A-Z:
                   
                   E. Daftar Pustaka
                   
                   Buku / Monograf / Book Chapter
                   [Daftar seluruh rujukan buku disajikan secara terurut alfabetis dari A sampai Z, format APA Style 7th]
                   
                   Jurnal Ilmiah / Artikel Jurnal
                   [Daftar seluruh rujukan artikel jurnal disajikan secara terurut alfabetis dari A sampai Z, format APA Style 7th dengan DOI]
                """
                
                response = generate_content_with_retry(
                    client,
                    selected_model,
                    [prompt_draft]
                )
                st.session_state['draft_result'] = response.text

    if st.session_state['draft_result']:
        st.markdown("---")
        st.subheader("📄 Draf Artikel Ilmiah Complete")
        st.markdown(st.session_state['draft_result'])
        
        st.markdown("---")
        st.subheader("💬 Ruang Diskusi & Penyesuaian Draf Naskah")
        st.caption("Berikan komentar, koreksi substansi, atau permintaan penambahan sitasi/referensi sebelum mengunduh naskah.")
        
        draft_comment = st.text_input("Catatan / Komentar Penyesuaian Draf:", placeholder="Misal: Tambahkan analisis perbandingan pada Bab Pembahasan dan perbanyak sitasi buku filsafat...")
        if st.button("🔄 Perbarui Draf Naskah Sesuai Catatan"):
            with st.spinner("Memperbarui dan merevisi draf naskah secara akademis..."):
                prompt_refine_draft = f"""Berikut Draf Naskah saat ini:\n{st.session_state['draft_result']}\n\nCatatan Tambahan Penulis:\n{draft_comment}\n\nTolong perbarui draf naskah tersebut secara komprehensif dengan tetap mematuhi Aturan Emas dan menyusun Daftar Pustaka terpisah (Buku A-Z & Jurnal A-Z) dalam format APA Style 7th Edition."""
                response_refine_draft = generate_content_with_retry(
                    client,
                    selected_model,
                    prompt_refine_draft
                )
                st.session_state['draft_result'] = response_refine_draft.text
                st.rerun()

        st.markdown("---")
        st.subheader("📥 Unduh Dokumen Naskah Utuh")
        docx_file = markdown_to_docx(st.session_state['draft_result'])
        st.download_button(
            label="📘 Unduh Naskah (.docx / Microsoft Word)",
            data=docx_file,
            file_name=f"Draf_{clean_proj_filename}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

# TAB 4: ASISTEN REVISI JURNAL (PEER REVIEW)
with tab4:
    st.header("🛠️ Asisten Pendamping Revisi Artikel & Peer Review")
    st.markdown("Fitur ini membantu menyusun **Matriks Tanggapan Reviewer (*Response to Reviewers*)** dan **Draf Revisi Naskah** secara santun, akademis, dan presisi.")
    
    col1, col2 = st.columns(2)
    with col1:
        st.subheader("1. Unggah / Masukkan Naskah Awal")
        manuscript_file = st.file_uploader("Unggah Naskah Asli (PDF):", type=["pdf"], key="uploader_ms")
        manuscript_text = st.text_area("Atau Tempel Teks Naskah Asli di Sini:", height=150)
        
    with col2:
        st.subheader("2. Unggah / Masukkan Catatan Reviewer")
        reviewer_file = st.file_uploader("Unggah Catatan Reviewer / Editor (PDF):", type=["pdf"], key="uploader_rev")
        reviewer_text = st.text_area("Atau Tempel Catatan/Komentar Reviewer di Sini:", height=150)

    if st.button("🚀 Process Revision & Generate Response Matrix"):
        if (manuscript_file or manuscript_text) and (reviewer_file or reviewer_text):
            with st.spinner("AI sedang menganalisis catatan reviewer dan menyusun tanggapan revisi..."):
                try:
                    contents_payload = []
                    temp_paths = []
                    
                    if manuscript_file:
                        with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
                            tmp.write(manuscript_file.getvalue())
                            temp_paths.append(tmp.name)
                        g_ms = client.files.upload(file=tmp.name)
                        contents_payload.append(g_ms)
                    elif manuscript_text:
                        contents_payload.append(f"NASKAH ASLI PENGGUNA:\n{manuscript_text}")
                        
                    if reviewer_file:
                        with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
                            tmp.write(reviewer_file.getvalue())
                            temp_paths.append(tmp.name)
                        g_rev = client.files.upload(file=tmp.name)
                        contents_payload.append(g_rev)
                    elif reviewer_text:
                        contents_payload.append(f"CATATAN/KOMENTAR REVIEWER:\n{reviewer_text}")
                    
                    prompt_rev = SYSTEM_INSTRUCTION + """\n
                    Tugas: Analisis masukan dari Reviewer/Editor dan buatkan dua keluaran utama:
                    
                    BAGIAN 1: TABEL/MATRIKS TANGGAPAN REVISI (RESPONSE TO REVIEWERS MATRIX)
                    Sajikan dalam format tabel/poin berurutan:
                    - Poin Komentar Reviewer
                    - Tanggapan Penulis (Gunakan bahasa akademis yang sangat santun dan apresiatif)
                    - Tindakan Perbaikan (Penjelasan bagian mana pada naskah yang diubah/ditambah)

                    BAGIAN 2: DRAF NASKAH HASIL REVISI (REVISED MANUSCRIPT)
                    Susun ulang naskah artikel yang sudah disempurnakan sesuai seluruh catatan reviewer di atas.
                    """
                    
                    response_rev = generate_content_with_retry(
                        client,
                        selected_model,
                        [prompt_rev] + contents_payload
                    )
                    st.session_state['revision_matrix'] = response_rev.text
                    
                    for path in temp_paths:
                        if os.path.exists(path):
                            os.remove(path)
                            
                except Exception as e:
                    st.error(f"Gagal memproses revisi: {e}")
        else:
            st.warning("Mohon sediakan berkas/teks Naskah Asli DAN Catatan Reviewer terlebih dahulu.")

    if st.session_state['revision_matrix']:
        st.markdown("---")
        st.subheader("📋 Matriks Tanggapan & Draf Naskah Terrevisi")
        st.markdown(st.session_state['revision_matrix'])
        
        st.markdown("---")
        st.subheader("💬 Diskusi Interaktif Penyesuaian Revisi")
        revisi_comment = st.text_input("Ingin mengubah atau menambah argumen pada tanggapan reviewer tertentu?", placeholder="Misal: Tambahkan penjelasan pengujian sampel tambahan...")
        if st.button("🔄 Perbarui Revisi Sesuai Catatan"):
            with st.spinner("Memperbarui matriks & draf revisi..."):
                prompt_refine = f"Berikut hasil revisi saat ini:\n{st.session_state['revision_matrix']}\n\nMasukan Tambahan Penulis:\n{revisi_comment}\n\nTolong perbarui matriks dan draf naskah tersebut."
                response_refine = generate_content_with_retry(
                    client,
                    selected_model,
                    prompt_refine
                )
                st.session_state['revision_matrix'] = response_refine.text
                st.rerun()

        st.markdown("---")
        st.subheader("📥 Unduh Berkas Revisi")
        docx_rev = markdown_to_docx(st.session_state['revision_matrix'])
        st.download_button(
            label="📘 Unduh Matriks & Naskah Revisi (.docx / Word)",
            data=docx_rev,
            file_name=f"Revisi_{clean_proj_filename}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

# TAB 5: EKSPOR SITASI RIS (MENDELEY / ZOTERO)
with tab5:
    st.header("📚 Ekspor Sitasi RIS ke Mendeley / Zotero")
    st.markdown("Tab ini secara otomatis mengonversi seluruh rujukan **Buku, Monograf, dan Artikel Jurnal** dari Crossref, OpenAlex, dan Google Books menjadi berkas **`.ris`** siap pakai.")
    
    if st.session_state['hybrid_citations']:
        st.success(f"Terdeteksi **{len(st.session_state['hybrid_citations'])} rujukan resmi (Buku & Jurnal)** dari draf naskah saat ini.")
        
        ris_data_string = convert_citations_to_ris(st.session_state['hybrid_citations'])
        
        st.subheader("📑 Pratinjau Teks Kode RIS:")
        st.code(ris_data_string[:1000] + ("\n..." if len(ris_data_string) > 1000 else ""), language="text")
        
        st.markdown("---")
        st.download_button(
            label="📥 Unduh Berkas Sitasi (.ris) untuk Mendeley",
            data=ris_data_string,
            file_name=f"Sitasi_{clean_proj_filename}.ris",
            mime="application/x-research-info-systems",
            help="Impor berkas ini ke Mendeley: Add New -> Import Library -> RIS (.ris)"
        )
    else:
        st.info("💡 Belum ada data sitasi Hybrid. Silakan jalankan tombol **Generasi Draf Naskah Lengkap** di Tab 3 terlebih dahulu.")
